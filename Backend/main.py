from fastapi import FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, HTMLResponse
from pydantic import BaseModel, Field
from openrouter import OpenRouter
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import requests
try:
    from .database import (
        create_customer as db_create_customer,
        create_ticket as db_create_ticket,
        find_customer as db_find_customer,
        find_customer_by_phone as db_find_customer_by_phone,
        find_customer_by_query as db_find_customer_by_query,
        get_customer_history,
        get_connection,
        get_or_create_chatbot,
        create_new_chatbot,
        get_chatbot_by_token,
        get_knowledge_sources,
        save_knowledge_source,
        get_ticket as db_get_ticket,
        log_message,
        normalize_phone,
        update_ticket as db_update_ticket,
    )
except ImportError:
    from database import (
        create_customer as db_create_customer,
        create_ticket as db_create_ticket,
        find_customer as db_find_customer,
        find_customer_by_phone as db_find_customer_by_phone,
        find_customer_by_query as db_find_customer_by_query,
        get_customer_history,
        get_connection,
        get_or_create_chatbot,
        create_new_chatbot,
        get_chatbot_by_token,
        get_knowledge_sources,
        save_knowledge_source,
        get_ticket as db_get_ticket,
        log_message,
        normalize_phone,
        update_ticket as db_update_ticket,
    )
import os
import re
import uuid
from html.parser import HTMLParser
from urllib.request import Request as UrlRequest, urlopen
from urllib.parse import urldefrag, urljoin, urlparse
from collections import deque

# =========================================================
# ENVIRONMENT
# =========================================================
load_dotenv()
API_KEY = os.getenv("OPENROUTER_API_KEY")

if not API_KEY:
    print("WARNING: OPENROUTER_API_KEY not found in .env")

# =========================================================
# FASTAPI
# =========================================================
app = FastAPI(title="NetVeda AI", version="10.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# =========================================================
# PROJECT PATHS
# =========================================================
BASE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = BASE_DIR.parent
DATA_DIR = PROJECT_DIR / "Data"
KNOWLEDGE_DIR = PROJECT_DIR / "Knowledge"
FRONTEND_DIR = PROJECT_DIR / "Frontend"

DATA_DIR.mkdir(exist_ok=True)
KNOWLEDGE_DIR.mkdir(exist_ok=True)

LEADS_FILE = DATA_DIR / "leads.xlsx"
SUPPORT_FILE = DATA_DIR / "support_tickets.xlsx"
ORDERS_FILE = DATA_DIR / "orders_subscriptions.xlsx"
KNOWLEDGE_FILE = KNOWLEDGE_DIR / "netveda.txt"
CUSTOMER_RECORDS_FILE = DATA_DIR / "customer_ticket_records.xlsx"

print("=" * 50)
print("NETVEDA BACKEND")
print("=" * 50)
print("BASE DIR:", BASE_DIR)
print("PROJECT DIR:", PROJECT_DIR)
print("DATA DIR:", DATA_DIR)
print("LEADS FILE:", LEADS_FILE)
print("SUPPORT FILE:", SUPPORT_FILE)
print("ORDERS FILE:", ORDERS_FILE)
print("KNOWLEDGE FILE:", KNOWLEDGE_FILE)
print("FRONTEND DIR:", FRONTEND_DIR)
print("=" * 50)

# =========================================================
# OPENROUTER
# =========================================================
client = None

if API_KEY:
    try:
        client = OpenRouter(api_key=API_KEY)
        print("OpenRouter client ready.")
    except Exception as e:
        print("OpenRouter initialization error:", e)

# =========================================================
# TIME
# =========================================================
def current_time():
    return datetime.now(ZoneInfo("Asia/Kolkata")).strftime(
        "%Y-%m-%d %H:%M:%S"
    )

# =========================================================
# EXCEL - LEADS
# EXACT COLUMNS ONLY:
# Name | Email | Phone | Requirement | Interested Plan |
# Message | Date & Time
# =========================================================
LEAD_HEADERS = [
    "Name",
    "Email",
    "Phone",
    "Requirement",
    "Interested Plan",
    "Message",
    "Date & Time",
]


def create_leads_file():
    if not LEADS_FILE.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Customer Leads"
        ws.append(LEAD_HEADERS)
        wb.save(LEADS_FILE)
        print("Created leads.xlsx")


def save_lead(data):
    create_leads_file()
    wb = load_workbook(LEADS_FILE)
    ws = wb["Customer Leads"]

    ws.append([
        data.get("name", ""),
        data.get("email", ""),
        data.get("phone", ""),
        data.get("requirement", ""),
        data.get("interested_plan", ""),
        data.get("message", ""),
        current_time(),
    ])

    wb.save(LEADS_FILE)
    print("Lead saved to leads.xlsx")

# =========================================================
# EXCEL - SUPPORT TICKETS
# EXACT COLUMNS ONLY:
# Ticket ID | Customer ID | Issue | Area | Priority |
# Status | Location
# =========================================================
SUPPORT_HEADERS = [
    "Ticket ID",
    "Customer ID",
    "Issue",
    "Area",
    "Priority",
    "Status",
    "Location",
]


def create_support_file():
    if not SUPPORT_FILE.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Support Tickets"
        ws.append(SUPPORT_HEADERS)
        wb.save(SUPPORT_FILE)
        print("Created support_tickets.xlsx")


def generate_ticket_id():
    return "TKT-" + uuid.uuid4().hex[:8].upper()


def create_support_ticket(data):
    create_support_file()
    ticket_id = generate_ticket_id()

    wb = load_workbook(SUPPORT_FILE)
    ws = wb["Support Tickets"]
    ws.append([
        ticket_id,
        data.get("customer_id", ""),
        data.get("issue", ""),
        data.get("area", ""),
        data.get("priority", "Medium"),
        "Open",
        data.get("location", ""),
    ])
    wb.save(SUPPORT_FILE)

    print("Support ticket saved:", ticket_id)
    return ticket_id

# =========================================================
# EXCEL - ORDERS & SUBSCRIPTIONS
# EXACT COLUMNS:
# Order ID | Customer ID | Name | Phone | Plan | Amount |
# Action | Status | Date & Time
# =========================================================
ORDER_HEADERS = [
    "Order ID",
    "Customer ID",
    "Name",
    "Phone",
    "Plan",
    "Amount",
    "Action",
    "Status",
    "Date & Time",
]

CUSTOMER_RECORD_HEADERS = [
    "Customer ID",
    "Name",
    "Phone Number",
    "Email",
    "Product",
    "Source Surface",
    "Ticket ID",
    "Issue / Query",
    "Category",
    "Priority",
    "Ticket Status",
    "Handled By",
    "Created At",
    "Updated At",
]


def save_customer_ticket_record(customer, ticket=None):
    if not CUSTOMER_RECORDS_FILE.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Customer Ticket Records"
        sheet.append(CUSTOMER_RECORD_HEADERS)
        workbook.save(CUSTOMER_RECORDS_FILE)

    workbook = load_workbook(CUSTOMER_RECORDS_FILE)
    sheet = workbook["Customer Ticket Records"]

    existing_ticket_ids = {
        str(row[6].value)
        for row in sheet.iter_rows(min_row=2)
        if row[6].value
    }
    if ticket and ticket["ticket_id"] in existing_ticket_ids:
        return

    sheet.append([
        customer["customer_id"],
        customer["name"],
        customer["phone_number"],
        customer.get("email"),
        customer.get("product"),
        ticket["source_surface"] if ticket else customer["source_surface"],
        ticket["ticket_id"] if ticket else "",
        ticket["query_text"] if ticket else "",
        ticket["category"] if ticket else "",
        ticket["priority"] if ticket else "",
        ticket["status"] if ticket else "",
        ticket["handled_by"] if ticket else "",
        ticket["created_at"] if ticket else customer["created_at"],
        ticket["updated_at"] if ticket else customer["last_active_at"],
    ])
    workbook.save(CUSTOMER_RECORDS_FILE)


def create_orders_file():
    if not ORDERS_FILE.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Orders & Subscriptions"
        ws.append(ORDER_HEADERS)
        wb.save(ORDERS_FILE)
        print("Created orders_subscriptions.xlsx")


def generate_order_id():
    return "ORD-" + uuid.uuid4().hex[:8].upper()


def create_order_subscription(data):
    create_orders_file()
    order_id = generate_order_id()

    wb = load_workbook(ORDERS_FILE)
    ws = wb["Orders & Subscriptions"]
    ws.append([
        order_id,
        data.get("customer_id", ""),
        data.get("name", ""),
        data.get("phone", ""),
        data.get("plan", ""),
        data.get("amount", ""),
        data.get("action", "New Subscription"),
        data.get("status", "Pending"),
        current_time(),
    ])
    wb.save(ORDERS_FILE)
    return order_id


def plan_amount(plan):
    return {
        "NetVeda 199": 199,
        "NetVeda 299": 299,
        "NetVeda 399": 399,
    }.get(plan, "")

# =========================================================
# CREATE ALL THREE FILES AT STARTUP
# =========================================================
create_leads_file()
create_support_file()
create_orders_file()
print("EXCEL FILES READY")
print("LEADS:", LEADS_FILE)
print("SUPPORT:", SUPPORT_FILE)
print("ORDERS:", ORDERS_FILE)

# =========================================================
# CUSTOMER ID
# =========================================================
def generate_customer_id():
    return "NV" + uuid.uuid4().hex[:6].upper()

# =========================================================
# CUSTOMER SEARCH FROM LEADS
# leads.xlsx intentionally has NO Customer ID column.
# =========================================================
def find_customer_by_phone(phone):
    if not LEADS_FILE.exists():
        return []

    try:
        wb = load_workbook(LEADS_FILE, data_only=True)
        ws = wb["Customer Leads"]
    except Exception as e:
        print("Lead read error:", e)
        return []

    phone = re.sub(r"\D", "", str(phone))
    results = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 7:
            continue

        stored_phone = re.sub(r"\D", "", str(row[2] or ""))
        if stored_phone and stored_phone == phone:
            results.append({
                "name": row[0] or "",
                "email": row[1] or "",
                "phone": stored_phone,
                "requirement": row[3] or "",
                "plan": row[4] or "",
                "message": row[5] or "",
            })

    return results

# =========================================================
# CUSTOMER ID LOOKUP
# Customer IDs are stored in orders/support because leads.xlsx
# is intentionally restricted to the 7 requested columns.
# =========================================================
def customer_id_exists(customer_id):
    customer_id = str(customer_id).strip()
    if not customer_id:
        return False

    for file_path, sheet_name, column_index in [
        (SUPPORT_FILE, "Support Tickets", 1),
        (ORDERS_FILE, "Orders & Subscriptions", 1),
    ]:
        if not file_path.exists():
            continue

        try:
            wb = load_workbook(file_path, data_only=True)
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and len(row) > column_index:
                    if str(row[column_index] or "").strip() == customer_id:
                        return True
        except Exception as e:
            print("Customer ID lookup error:", e)

    return False

# =========================================================
# VERIFY EXISTING CUSTOMER TICKET
# Phone is verified against leads.xlsx and the customer/ticket
# pair is verified against support_tickets.xlsx.
# =========================================================
def verify_existing_customer(phone, customer_id, ticket_id):
    if not find_customer_by_phone(phone):
        return False

    if not SUPPORT_FILE.exists():
        return False

    try:
        wb = load_workbook(SUPPORT_FILE, data_only=True)
        ws = wb["Support Tickets"]
    except Exception:
        return False

    customer_id = str(customer_id).strip()
    ticket_id = str(ticket_id).strip()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 7:
            continue

        if (
            str(row[0] or "").strip() == ticket_id
            and str(row[1] or "").strip() == customer_id
        ):
            return True

    return False

# =========================================================
# GET CUSTOMER TICKETS
# =========================================================
def get_customer_tickets(customer_id):
    if not SUPPORT_FILE.exists():
        return []

    try:
        wb = load_workbook(SUPPORT_FILE, data_only=True)
        ws = wb["Support Tickets"]
    except Exception:
        return []

    tickets = []
    customer_id = str(customer_id).strip()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 7:
            continue

        if str(row[1] or "").strip() == customer_id:
            tickets.append({
                "ticket_id": row[0],
                "customer_id": row[1],
                "issue": row[2],
                "area": row[3],
                "priority": row[4],
                "status": row[5],
                "location": row[6],
            })

    return tickets

# =========================================================
# KNOWLEDGE BASE
# =========================================================
def load_knowledge():
    if not KNOWLEDGE_FILE.exists():
        return ""
    try:
        return KNOWLEDGE_FILE.read_text(encoding="utf-8")
    except Exception as e:
        print("Knowledge loading error:", e)
        return ""


KNOWLEDGE_TEXT = load_knowledge()
print("Knowledge characters loaded:", len(KNOWLEDGE_TEXT))

knowledge_chunks = []
if KNOWLEDGE_TEXT:
    knowledge_chunks = [
        x.strip()
        for x in re.split(r"\n\s*\n", KNOWLEDGE_TEXT)
        if x.strip()
    ]

vectorizer = None
knowledge_vectors = None

if knowledge_chunks:
    try:
        vectorizer = TfidfVectorizer(stop_words="english")
        knowledge_vectors = vectorizer.fit_transform(knowledge_chunks)
        print("Knowledge search system ready.")
    except Exception as e:
        print("Knowledge vectorization error:", e)


def search_knowledge(query, top_k=3):
    if not knowledge_chunks or vectorizer is None or knowledge_vectors is None:
        return ""

    try:
        query_vector = vectorizer.transform([query])
        similarities = cosine_similarity(
            query_vector,
            knowledge_vectors
        )[0]
        indexes = similarities.argsort()[::-1][:top_k]

        results = []
        for index in indexes:
            if similarities[index] > 0:
                results.append(knowledge_chunks[index])

        return "\n\n".join(results)
    except Exception:
        return ""

# =========================================================
# SESSION
# NOTE: This is an in-memory session only. Normal chat is NOT
# written to Excel and no conversation history is stored.
# =========================================================
customer_session = {}


def reset_session():
    customer_session.clear()
    customer_session.update({
        "mode": None,
        "step": None,
        "customer_id": None,
        "name": None,
        "email": None,
        "phone": None,
        "requirement": None,
        "interested_plan": None,
        "message": None,
        "area": None,
        "location": None,
        "issue": None,
        "priority": "Medium",
        "original_message": None,
        "existing_customer": False,
        "action": None,
    })


reset_session()

# =========================================================
# TOPIC DETECTION
# =========================================================
def is_existing_customer_request(message):
    text = message.lower()
    words = [
        "existing customer",
        "old customer",
        "my old connection",
        "previous ticket",
        "my previous ticket",
        "my connection",
        "already a customer",
        "i am already a customer",
        "existing connection",
        "upgrade my plan",
        "downgrade my plan",
        "upgrade plan",
        "downgrade plan",
        "recharge my plan",
    ]
    return any(word in text for word in words)


def is_support_issue(message):
    text = message.lower()
    keywords = [
        "not working",
        "doesn't work",
        "doesnt work",
        "internet down",
        "internet not working",
        "wifi not working",
        "wi-fi not working",
        "wifi down",
        "slow internet",
        "network issue",
        "network problem",
        "connection problem",
        "connection issue",
        "no internet",
        "no signal",
        "signal problem",
        "router problem",
        "service problem",
        "service issue",
        "complaint",
        "issue",
        "problem",
        "support ticket",
    ]
    return any(word in text for word in keywords)


def is_order_request(message):
    text = message.lower()
    keywords = [
        "buy",
        "purchase",
        "subscribe",
        "subscription",
        "order",
        "new connection",
        "take this plan",
        "get this plan",
        "activate plan",
        "upgrade",
        "downgrade",
        "recharge",
    ]
    return any(word in text for word in keywords)


def is_plan_request(message):
    text = message.lower()
    keywords = [
        "plan",
        "recharge",
        "data plan",
        "gaming",
        "internet plan",
        "best plan",
        "which plan",
        "recommend plan",
        "data requirement",
    ]
    return any(word in text for word in keywords)

# =========================================================
# PLAN HELPERS
# =========================================================
def detect_plan(message):
    text = message.lower()
    if "199" in text:
        return "NetVeda 199"
    if "299" in text:
        return "NetVeda 299"
    if "399" in text:
        return "NetVeda 399"
    return None


def recommend_plan(message, company_name="NetVeda"):
    text = message.lower()

    if "gaming" in text or "game" in text or "gamer" in text:
        return (
            f"For gaming, I recommend **{company_name} 399**.\n\n"
            "Price: INR 399\n"
            "Data: 2 GB/day\n"
            "Validity: 56 days\n\n"
            f"This matches the current {company_name} gaming-focused requirement."
        )

    return (
        f"Here are the available {company_name} plans:\n\n"
        f"**{company_name} 199**\n"
        "- INR 199\n"
        "- 1 GB/day\n"
        "- 28 days\n\n"
        f"**{company_name} 299**\n"
        "- INR 299\n"
        "- 1.5 GB/day\n"
        "- 28 days\n\n"
        f"**{company_name} 399**\n"
        "- INR 399\n"
        "- 2 GB/day\n"
        "- 56 days"
    )

# =========================================================
# START NEW CUSTOMER FLOW
# =========================================================
def start_new_customer_flow(message, action="New Connection"):
    reset_session()

    customer_session["mode"] = "new_customer"
    customer_session["step"] = "name"
    customer_session["customer_id"] = generate_customer_id()
    customer_session["original_message"] = message
    customer_session["message"] = message
    customer_session["action"] = action

    detected = detect_plan(message)
    if detected:
        customer_session["interested_plan"] = detected

    return (
        "Sure! I'll help you with your NetVeda request. 😊\n\n"
        "First, please provide your **full name**."
    )

# =========================================================
# START EXISTING CUSTOMER FLOW
# =========================================================
def start_existing_customer_flow(message):
    reset_session()
    customer_session["mode"] = "existing_verification"
    customer_session["step"] = "phone"
    customer_session["original_message"] = message

    text = message.lower()
    if "upgrade" in text:
        customer_session["action"] = "Upgrade"
    elif "downgrade" in text:
        customer_session["action"] = "Downgrade"
    elif "recharge" in text:
        customer_session["action"] = "Recharge"
    else:
        customer_session["action"] = "Support"

    return (
        "Sure. Since this is an existing customer request, "
        "I first need to verify your connection.\n\n"
        "Please provide your **registered phone number**."
    )

# =========================================================
# NEW CUSTOMER FLOW
# =========================================================
def handle_new_customer(message):
    step = customer_session.get("step")

    # NAME
    if step == "name":
        customer_session["name"] = message
        customer_session["step"] = "email"
        return (
            "Thank you. 👍\n\n"
            "Please provide your **email address**."
        )

    # EMAIL
    if step == "email":
        if "@" not in message or "." not in message.split("@")[-1]:
            return "Please provide a valid email address."

        customer_session["email"] = message
        customer_session["step"] = "phone"
        return (
            "Great. Now please provide your **registered phone number**."
        )

    # PHONE
    if step == "phone":
        phone = re.sub(r"\D", "", message)
        if len(phone) < 10:
            return "Please provide a valid 10-digit phone number."

        customer_session["phone"] = phone
        customer_session["step"] = "requirement"
        return (
            "Thank you. 👍\n\n"
            "Please tell me your **requirement**."
        )

    # REQUIREMENT
    if step == "requirement":
        customer_session["requirement"] = message

        if (
            "gaming" in message.lower()
            or "game" in message.lower()
            or "gamer" in message.lower()
        ):
            customer_session["interested_plan"] = "NetVeda 399"
            customer_session["step"] = "message"
            return (
                "🎮 Based on your gaming requirement, "
                "I recommend **NetVeda 399**.\n\n"
                "💰 ₹399\n"
                "📶 2 GB/day\n"
                "📅 56 days\n\n"
                "Now please tell me any additional message or "
                "details you want to share."
            )

        if customer_session.get("interested_plan"):
            customer_session["step"] = "message"
            return (
                "Got it. 👍\n\n"
                f"Selected Plan: **{customer_session['interested_plan']}**\n\n"
                "Please share any **additional message or details**."
            )

        customer_session["step"] = "interested_plan"
        return (
            "Got it. 👍\n\n"
            "Which NetVeda plan are you interested in?\n\n"
            "You can choose **199, 299, 399**, or simply say **Not sure**."
        )

    # PLAN
    if step == "interested_plan":
        plan = detect_plan(message)

        if not plan and "not sure" in message.lower():
            requirement = customer_session.get("requirement", "").lower()
            plan = "NetVeda 399" if ("gaming" in requirement or "game" in requirement) else "Not decided"

        if not plan:
            return (
                "Please select NetVeda 199, 299, 399, "
                "or say **Not sure**."
            )

        customer_session["interested_plan"] = plan
        customer_session["step"] = "message"
        return (
            "Perfect. 👍\n\n"
            "Please share any **additional message or details** you want to provide."
        )

    # MESSAGE
    if step == "message":
        customer_session["message"] = message
        customer_session["step"] = "area"
        return (
            "Thank you. 👍\n\n"
            "Please provide your **area**."
        )

    # AREA
    if step == "area":
        customer_session["area"] = message
        customer_session["step"] = "location"
        return (
            "Thanks. Now please provide your **location/city**."
        )

    # LOCATION
    if step == "location":
        customer_session["location"] = message
        action = customer_session.get("action", "New Connection")

        # SUPPORT: DO NOT SAVE TO LEADS
        if is_support_issue(customer_session.get("original_message", "")):
            customer_session["step"] = "issue"
            return (
                "Thank you. I have your customer details.\n\n"
                "Now describe the **issue** you are facing."
            )

        # ORDER / SUBSCRIPTION: SAVE ONLY TO ORDERS FILE
        if action in {
            "New Subscription",
            "Upgrade",
            "Downgrade",
            "Recharge",
        }:
            return create_order_response()

        # NEW CONNECTION / LEAD: SAVE ONLY TO LEADS FILE
        save_lead(customer_session)

        response = (
            "✅ Your customer details have been registered successfully.\n\n"
            f"👤 Customer ID: {customer_session['customer_id']}\n"
        )

        if customer_session.get("interested_plan"):
            response += (
                f"📱 Interested Plan: {customer_session['interested_plan']}\n"
            )

        response += "\nThank you for choosing NetVeda! 😊"
        reset_session()
        return response

    # ISSUE
    # IMPORTANT FIX: after receiving the issue, create the ticket
    # and RESET the session. This prevents the same troubleshooting
    # message from repeating for every next message such as "hi".
    if step == "issue":
        customer_session["issue"] = message
        return create_followup_ticket_response()

    return "I understand. Please continue with your request."

# =========================================================
# CREATE ORDER RESPONSE
# =========================================================
def create_order_response():
    plan = customer_session.get("interested_plan")

    if not plan or plan == "Not decided":
        plan = detect_plan(customer_session.get("original_message", ""))

    if not plan:
        plan = "Not decided"

    amount = plan_amount(plan)
    action = customer_session.get("action", "New Subscription")

    if action == "New Connection":
        action = "New Subscription"

    order_id = create_order_subscription({
        "customer_id": customer_session.get("customer_id", ""),
        "name": customer_session.get("name", ""),
        "phone": customer_session.get("phone", ""),
        "plan": plan,
        "amount": amount,
        "action": action,
        "status": "Pending",
    })

    response = (
        "🛒 **Order / Subscription Created Successfully!** ✅\n\n"
        f"🆔 Order ID: {order_id}\n"
        f"👤 Customer ID: {customer_session.get('customer_id', '')}\n"
        f"👤 Name: {customer_session.get('name', '')}\n"
        f"📱 Phone: {customer_session.get('phone', '')}\n"
        f"📦 Plan: {plan}\n"
    )

    if amount != "":
        response += f"💰 Amount: ₹{amount}\n"

    response += (
        f"🔄 Action: {action}\n"
        "📌 Status: Pending\n\n"
        "Your order has been recorded successfully. 😊"
    )

    reset_session()
    return response

# =========================================================
# AUTOMATIC SUPPORT TICKET RESPONSE
# =========================================================
def create_followup_ticket_response():
    ticket_id = create_support_ticket(customer_session)

    customer_id = customer_session.get("customer_id", "")
    issue = customer_session.get("issue", "Service issue")
    area = customer_session.get("area", "")
    location = customer_session.get("location", "")
    priority = customer_session.get("priority", "Medium")

    response = (
        "🎫 **Support Ticket Created Successfully!** ✅\n\n"
        f"👤 Customer ID: {customer_id}\n\n"
        f"🎫 Ticket ID: {ticket_id}\n\n"
        f"📋 Issue: {issue}\n\n"
        f"📍 Area: {area}\n\n"
        f"🌍 Location: {location}\n\n"
        "📌 Status: Open\n"
        f"⚡ Priority: {priority}\n\n"
        "Your query has been registered successfully. "
        "Thank you for your patience. 🙏"
    )

    reset_session()
    return response

# =========================================================
# EXISTING CUSTOMER VERIFICATION FLOW
# =========================================================
def handle_existing_verification(message):
    step = customer_session.get("step")

    # PHONE
    if step == "phone":
        phone = re.sub(r"\D", "", message)

        if len(phone) < 10:
            return "Please provide your registered 10-digit phone number."

        customer_session["phone"] = phone
        customers = find_customer_by_phone(phone)

        if not customers:
            # IMPORTANT: reset AND immediately start new-customer mode.
            reset_session()
            customer_session["mode"] = "new_customer"
            customer_session["step"] = "name"
            customer_session["customer_id"] = generate_customer_id()
            customer_session["original_message"] = message
            customer_session["action"] = "New Connection"

            return (
                "I couldn't find a NetVeda connection registered with this phone number.\n\n"
                "I'll register you as a new customer.\n\n"
                "Please provide your **full name**."
            )

        customer_session["existing_customer"] = True
        customer_session["existing_records"] = customers
        customer_session["step"] = "customer_id"

        return (
            "I found a NetVeda connection for this phone number. ✅\n\n"
            "Please provide your **Customer ID**."
        )

    # CUSTOMER ID
    if step == "customer_id":
        entered_customer_id = str(message).strip()

        if not customer_id_exists(entered_customer_id):
            return (
                "The Customer ID could not be verified.\n\n"
                "Please provide the correct Customer ID."
            )

        customer_session["customer_id"] = entered_customer_id
        records = customer_session.get("existing_records", [])

        if records:
            customer_session["name"] = records[0].get("name", "")
            customer_session["email"] = records[0].get("email", "")
            customer_session["phone"] = records[0].get("phone", "")

        action = customer_session.get("action", "Support")

        if action in {"Upgrade", "Downgrade", "Recharge"}:
            customer_session["step"] = "order_plan"
            return (
                "Customer ID verified. ✅\n\n"
                f"Please provide the **NetVeda plan** you want for {action.lower()}."
            )

        customer_session["step"] = "ticket_id"
        return (
            "Customer ID verified. ✅\n\n"
            "Now please provide your **Ticket ID**."
        )

    # ORDER PLAN FOR EXISTING CUSTOMER
    if step == "order_plan":
        plan = detect_plan(message)
        if not plan:
            return "Please select NetVeda 199, 299, or 399."

        customer_session["interested_plan"] = plan
        customer_session["step"] = "order_message"
        return (
            f"Perfect. **{plan}** selected. 👍\n\n"
            "Please share any additional message or details."
        )

    # ORDER MESSAGE
    if step == "order_message":
        customer_session["message"] = message
        return create_order_response()

    # TICKET ID
    if step == "ticket_id":
        ticket_id = message.strip()

        if not verify_existing_customer(
            customer_session.get("phone", ""),
            customer_session.get("customer_id", ""),
            ticket_id,
        ):
            return (
                "I couldn't verify that Ticket ID with your registered phone number "
                "and Customer ID.\n\n"
                "Please provide the correct Ticket ID."
            )

        customer_session["ticket_id"] = ticket_id
        customer_session["step"] = "connection"

        tickets = get_customer_tickets(customer_session["customer_id"])

        if len(tickets) > 1:
            return (
                "Your customer details are verified. ✅\n\n"
                "You have multiple support tickets.\n\n"
                "Which **connection or issue** are you currently facing?"
            )

        if len(tickets) == 1:
            customer_session["issue"] = tickets[0]["issue"]
            customer_session["area"] = tickets[0]["area"]
            customer_session["location"] = tickets[0]["location"]
            customer_session["priority"] = tickets[0]["priority"] or "Medium"
            return create_followup_ticket_response()

        return (
            "Your details are verified. ✅\n\n"
            "Please tell me which connection or service you are facing an issue with."
        )

    # CONNECTION
    if step == "connection":
        customer_session["issue"] = message
        customer_session["step"] = "existing_area"
        return (
            "Got it. Please confirm the **area** where you are facing this issue."
        )

    # AREA
    if step == "existing_area":
        customer_session["area"] = message
        customer_session["step"] = "existing_location"
        return "Thanks. Please confirm your **location/city**."

    # LOCATION
    if step == "existing_location":
        customer_session["location"] = message
        return create_followup_ticket_response()

    return "Your details are being verified. Please continue."

# =========================================================
# UNIVERSAL AI RESPONSE
# No conversation_history is stored.
# =========================================================
def generate_ai_response(message, conversation=None, knowledge_extra="", use_global_knowledge=True):
    if not client:
        return (
            "I'm currently unable to connect to the AI service. "
            "Please check the backend API key."
        )

    knowledge = search_knowledge(message) if use_global_knowledge else ""

    system_prompt = """
You are Kairo AI, an intelligent and highly capable assistant.

CRITICAL INSTRUCTIONS (MUST FOLLOW STRICTLY):
1. STRICT LANGUAGE MATCHING: ALWAYS respond in the exact language the user uses. If the user asks in English, you MUST reply ONLY in English. Never reply in Hindi or any other language unless the user explicitly speaks in that language or requests it. Do not assume or change languages on your own.
2. DO NOT DO THINGS UNPROMPTED: Only answer exactly what the user asks. Do not add unsolicited advice, and do not behave in a way that wasn't requested.
3. NEVER TELL THE USER TO VISIT THE WEBSITE: If the user asks for plans, pricing, or services, YOU MUST list out whatever information you have. NEVER say "I recommend checking their official website". If you know the website domain, you MUST generate a clickable markdown link (e.g., [View Plans](https://domain.com/plans) or [View Services](https://domain.com/services)). Even if you have to guess the exact path like /plans or /services, provide the link!

YOUR PRIMARY GOALS:
1. Answer questions based on the provided WEBSITE KNOWLEDGE if the question is about this specific website. When mentioning a specific feature, plan, or page from the website, provide the absolute URL link using Markdown.
2. UNIVERSAL KNOWLEDGE: If the user asks about ANY other topic, you must provide a correct, detailed, and helpful answer using your general AI knowledge.

ADAPTIVE COMMUNICATION STYLE:
- Keep your answers natural and adapt the length to the user's need. If the question is simple, give a short, direct answer. If the question requires an explanation, give a detailed, longer answer.
- Be concise when appropriate, and expansive when needed. Be professional and clear.

Do not mention: backend code, internal session logic, OpenRouter, or API keys.
"""

    if knowledge:
        system_prompt += "\n\nLEGACY KNOWLEDGE:\n" + knowledge
    if knowledge_extra:
        system_prompt += "\n\nWEBSITE KNOWLEDGE:\n" + knowledge_extra[:24000]

    messages = [
        {"role": "system", "content": system_prompt},
    ]
    for item in (conversation or [])[-8:]:
        if item.get("role") in {"user", "assistant"} and item.get("content"):
            messages.append({
                "role": item["role"],
                "content": item["content"],
            })
    messages.append({"role": "user", "content": message})

    try:
        response = client.chat.send(
            model="openrouter/free",
            messages=messages,
        )
        reply = response.choices[0].message.content or ""
        
        # Handle unclosed or closed think tags
        reply = re.sub(r'<think>.*?(?:</think>|$)', '', reply, flags=re.DOTALL | re.IGNORECASE)
        reply = re.sub(r'<tool_call>.*?(?:</tool_call>|$)', '', reply, flags=re.DOTALL | re.IGNORECASE)
        
        reply = re.sub(r'(?im)^[ \t]*User Safety:.*$', '', reply)
        reply = re.sub(r'(?im)^[ \t]*Response Safety:.*$', '', reply)
        reply = re.sub(r'\n{3,}', '\n\n', reply)
        reply = reply.strip()
        
        if not reply:
            return "I apologize, but I couldn't generate a proper response to that. Could you please try asking in a slightly different way?"
            
        return reply
    except Exception as e:
        print("AI ERROR:", e)
        return (
            "Sorry, the AI service is temporarily unavailable or busy. "
            "Please try again in a few moments."
        )

# =========================================================
# REQUEST MODELS
# =========================================================
class ChatRequest(BaseModel):
    message: str


class TicketRequest(BaseModel):
    customer_id: str
    issue: str
    area: str
    priority: str = "Medium"
    location: str


class CustomerCreateRequest(BaseModel):
    name: str
    phone_number: str
    email: str | None = None
    product: str | None = None
    source_surface: str = "website"


class UnifiedTicketRequest(BaseModel):
    customer_id: str
    query_text: str
    category: str = "general"
    source_surface: str = "website"
    priority: str = "medium"


class TicketUpdateRequest(BaseModel):
    status: str | None = None
    resolution_notes: str | None = None
    priority: str | None = None
    handled_by: str | None = None


class ChatbotCreateRequest(BaseModel):
    name: str = "NetVeda AI"


class KnowledgeLinkRequest(BaseModel):
    url: str
    max_pages: int = Field(default=25, ge=1, le=100)


class KnowledgeTextRequest(BaseModel):
    text: str
    title: str = "Manual knowledge"


class VisibleTextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript", "svg"}:
            self.skip_depth += 1

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript", "svg"} and self.skip_depth:
            self.skip_depth -= 1

    def handle_data(self, data):
        if not self.skip_depth and data.strip():
            self.parts.append(data.strip())


class PageMetadataParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.title_parts = []
        self.description = ""
        self.in_title = False

    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        if tag == "title":
            self.in_title = True
        if tag == "meta" and attributes.get("name", "").lower() == "description":
            self.description = attributes.get("content", "").strip()

    def handle_endtag(self, tag):
        if tag == "title":
            self.in_title = False

    def handle_data(self, data):
        if self.in_title and data.strip():
            self.title_parts.append(data.strip())


class PageLinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []

    def handle_starttag(self, tag, attrs):
        if tag != "a":
            return
        href = dict(attrs).get("href", "").strip()
        if href:
            self.links.append(href)


def extract_visible_text(markup: str) -> str:
    parser = VisibleTextParser()
    parser.feed(markup)
    return "\n".join(parser.parts)


def extract_page_knowledge(markup: str) -> str:
    visible_text = extract_visible_text(markup)
    if len(visible_text) >= 20:
        return visible_text

    metadata_parser = PageMetadataParser()
    metadata_parser.feed(markup)
    metadata = []
    if metadata_parser.title_parts:
        metadata.append("Page title: " + " ".join(metadata_parser.title_parts))
    if metadata_parser.description:
        metadata.append("Page description: " + metadata_parser.description)
    return "\n".join(metadata + ([visible_text] if visible_text else []))


def extract_page_links(markup: str, page_url: str, hostname: str) -> list[str]:
    parser = PageLinkParser()
    parser.feed(markup)
    links = []
    for href in parser.links:
        absolute_url, _ = urldefrag(urljoin(page_url, href))
        parsed = urlparse(absolute_url)
        if parsed.scheme not in {"http", "https"} or parsed.hostname != hostname:
            continue
        if parsed.path.lower().endswith((".pdf", ".jpg", ".jpeg", ".png", ".gif", ".zip", ".mp4")):
            continue
        if absolute_url not in links:
            links.append(absolute_url)
    return links


def fetch_website_markup(url: str) -> str:
    response = requests.get(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (compatible; KairoKnowledgeBot/1.0)",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.8",
        },
        timeout=(3, 5),
        allow_redirects=True,
    )
    response.raise_for_status()
    content_type = response.headers.get("content-type", "").lower()
    if content_type and not any(value in content_type for value in ("text/html", "application/xhtml+xml", "text/plain")):
        raise ValueError("The URL did not return an HTML or text page")
    response.encoding = response.encoding or "utf-8"
    return response.text[:2_000_000]


def chatbot_embed_markup(public_token: str) -> str:
    return f"""<!doctype html>
<html lang="en">
<head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>NetVeda AI</title>
<style>body{{margin:0;font-family:Arial,sans-serif;background:#fff}}#app{{height:100vh}}</style></head>
<body><div id="app"></div><script>window.NETVEDA_CHATBOT_TOKEN={public_token!r};</script>
<script src="/embed.js"></script></body></html>"""


def chatbot_training_text(public_token: str | None) -> str:
    if not public_token:
        return ""
    chatbot = get_chatbot_by_token(public_token)
    if not chatbot:
        return ""
    return "\n\n".join(
        source["content"]
        for source in get_knowledge_sources(chatbot["chatbot_id"])
        if source.get("status") == "trained"
    )


class UnifiedChatRequest(BaseModel):
    message: str
    surface: str = "website"
    phone_number: str | None = None
    name: str | None = None
    email: str | None = None
    product: str | None = None
    employee_id: str | None = None
    customer_id: str | None = None
    customer_type: str | None = None
    conversation: list[dict[str, str]] = Field(default_factory=list)
    chatbot_token: str | None = None


def is_guest_eligible(message: str) -> bool:
    return is_plan_request(message) and not is_support_issue(message)


def classify_ticket(message: str) -> str:
    text = message.lower()
    if is_support_issue(message):
        return "support"
    if is_order_request(message):
        return "sales"
    if "demo" in text:
        return "demo_request"
    if "callback" in text:
        return "callback_request"
    return "general"


def is_greeting(message: str) -> bool:
    words = re.findall(r"[a-z]+", message.lower())
    return bool(words) and len(words) <= 5 and any(
        greeting in words
        for greeting in {"hi", "hello", "hey", "morning", "afternoon", "evening"}
    )


def wants_ticket_or_contact(message: str) -> bool:
    text = message.lower()
    phrases = [
        "raise a ticket",
        "create a ticket",
        "create a support ticket",
        "support ticket",
        "register a complaint",
        "contact support",
        "call me",
        "request a callback",
        "talk to an agent",
        "human agent",
        "new connection",
        "buy a plan",
        "subscribe",
        "recharge",
        "upgrade my plan",
        "downgrade my plan",
    ]
    return any(phrase in text for phrase in phrases)


def is_account_action(message: str) -> bool:
    text = message.lower()
    phrases = [
        "upgrade",
        "downgrade",
        "recharge",
        "subscribe",
        "purchase",
        "buy this plan",
        "get this plan",
        "activate this plan",
        "new connection",
    ]
    return any(phrase in text for phrase in phrases)


def conversational_reply(message: str) -> str:
    text = message.lower().strip()
    if is_greeting(message):
        return (
            "Hi! I’m Kairo AI. I can answer questions about this website or chat about general topics. "
            "What would you like to explore today?"
        )
    if text in {"thanks", "thank you", "thx"}:
        return "You’re welcome! Is there anything else you’d like help with?"
    if is_account_action(message):
        selected_plan = detect_plan(message)
        plan_text = f" ({selected_plan})" if selected_plan else ""
        return (
            f"I can help with that{plan_text}. Is this for an existing NetVeda customer, "
            "or should I help create a new connection? Reply **existing** or **new**."
        )
    if is_plan_request(message):
        return (
            f"{recommend_plan(message)}\n\n"
            "You can say **199**, **299**, or **399** to compare one plan. "
            "To buy, upgrade, or recharge, tell me what you want to do."
        )
    if "not working" in text or "no internet" in text or "no signal" in text:
        return (
            "I can help troubleshoot that. Please check whether the issue affects one device "
            "or all devices, and try restarting the router or switching mobile data off and on. "
            "Would you like me to register this as a support ticket if it continues?"
        )
    if "slow" in text and ("internet" in text or "network" in text or "wifi" in text):
        return (
            "For slow internet, try moving closer to the router, disconnecting unused devices, "
            "and restarting the connection. Does the slowdown happen on one device or all of them?"
        )
    if client:
        return generate_ai_response(message)
    return (
        "I’m happy to help with this website or with a normal conversation. What would you like to know?"
    )


# =========================================================
# UNIFIED CUSTOMER AND TICKET APIs
# =========================================================
@app.get("/customer/check")
def check_customer(phone: str):
    normalized = normalize_phone(phone)
    customer = db_find_customer_by_phone(normalized)
    return {
        "exists": customer is not None,
        "phone_number": normalized,
        "customer_id": customer["customer_id"] if customer else None,
        "customer": customer,
    }


@app.get("/customer/search")
def search_customers(query: str):
    return {"customers": db_find_customer_by_query(query)}


@app.get("/customer-records")
def customer_records():
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT
                c.customer_id, c.name, c.phone_number, c.email, c.product,
                c.source_surface, c.created_at AS customer_created_at,
                t.ticket_id, t.query_text, t.category, t.priority,
                t.status AS ticket_status, t.handled_by,
                t.created_at AS ticket_created_at, t.updated_at AS ticket_updated_at
            FROM customers c
            LEFT JOIN tickets t ON t.customer_id = c.customer_id
            ORDER BY c.created_at DESC, t.created_at DESC
            """
        ).fetchall()
    return {"records": [dict(row) for row in rows]}


@app.get("/customer/{customer_id}")
def customer_profile(customer_id: str):
    customer = db_find_customer(customer_id)
    if not customer:
        raise HTTPException(status_code=404, detail="Customer not found")
    return {
        "customer": customer,
        "history": get_customer_history(customer_id),
    }


@app.post("/customer/create", status_code=201)
def create_customer_api(request: CustomerCreateRequest):
    normalized = normalize_phone(request.phone_number)
    if len(normalized) < 10:
        raise HTTPException(status_code=422, detail="A valid phone number is required")

    existing = db_find_customer_by_phone(normalized)
    if existing:
        save_customer_ticket_record(existing)
        return {"created": False, "customer": existing}

    try:
        customer = db_create_customer(
            name=request.name,
            phone=normalized,
            email=request.email,
            product=request.product,
            source_surface=request.source_surface,
        )
    except Exception as error:
        raise HTTPException(status_code=409, detail="Customer could not be created") from error
    save_customer_ticket_record(customer)
    return {"created": True, "customer": customer}


@app.post("/ticket/create", status_code=201)
def create_unified_ticket(request: UnifiedTicketRequest):
    if not db_find_customer(request.customer_id):
        raise HTTPException(status_code=404, detail="Customer not found")
    ticket = db_create_ticket(
        customer_id=request.customer_id,
        query_text=request.query_text,
        category=request.category,
        source_surface=request.source_surface,
        priority=request.priority,
    )
    return {"ticket": ticket}


@app.get("/ticket/{ticket_id}")
def ticket_details(ticket_id: str):
    ticket = db_get_ticket(ticket_id)
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    return {"ticket": ticket}


@app.patch("/ticket/{ticket_id}")
def update_ticket_api(ticket_id: str, request: TicketUpdateRequest):
    if not db_get_ticket(ticket_id):
        raise HTTPException(status_code=404, detail="Ticket not found")
    ticket = db_update_ticket(ticket_id, request.model_dump(exclude_none=True))
    return {"ticket": ticket}


@app.post("/escalate")
def escalate_ticket(ticket_id: str, resolution_notes: str = "Escalated to a human agent"):
    ticket = db_update_ticket(
        ticket_id,
        {"status": "escalated", "resolution_notes": resolution_notes},
    )
    if not ticket:
        raise HTTPException(status_code=404, detail="Ticket not found")
    return {"ticket": ticket, "escalated": True}


@app.post("/chatbot/create", status_code=201)
def create_chatbot(request: ChatbotCreateRequest, http_request: Request):
    chatbot = create_new_chatbot(request.name)
    base_url = str(http_request.base_url).rstrip("/")
    return {
        "chatbot": chatbot,
        "public_url": f"{base_url}/chatbot/{chatbot['public_token']}",
        "embed_code": f'<script src="{base_url}/embed.js" data-chatbot-token="{chatbot["public_token"]}"></script>',
    }


@app.post("/chatbot/{public_token}/knowledge/link")
def train_from_link(public_token: str, request: KnowledgeLinkRequest):
    chatbot = get_chatbot_by_token(public_token)
    if not chatbot:
        raise HTTPException(status_code=404, detail="Chatbot not found")
    parsed_url = urlparse(request.url.strip())
    if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
        raise HTTPException(status_code=422, detail="A valid HTTP(S) URL is required")

    start_url = request.url.strip().rstrip("/")
    hostname = parsed_url.hostname
    queue = deque([start_url])
    visited = set()
    pages = []
    failures = []

    while queue and len(pages) < request.max_pages:
        page_url = queue.popleft()
        if page_url in visited:
            continue
        visited.add(page_url)
        try:
            markup = fetch_website_markup(page_url)
            content = extract_page_knowledge(markup)
            if len(content) >= 20:
                source = save_knowledge_source(chatbot["chatbot_id"], "website", page_url, content)
                pages.append({"url": page_url, "source": source})
            for discovered_url in extract_page_links(markup, page_url, hostname):
                if discovered_url not in visited and len(visited) < request.max_pages * 3:
                    queue.append(discovered_url)
        except Exception as error:
            print(f"Knowledge fetch failed for {page_url!r}: {error}")
            failures.append({"url": page_url, "error": str(error)})

    if not pages:
        detail = "Could not read any public pages from this website. Try its homepage, FAQ/help page, or paste the content manually."
        if failures and isinstance(failures[0].get("error"), str):
            detail += " Check that the URL is public and does not require login."
        raise HTTPException(status_code=422, detail=detail)

    return {
        "trained": True,
        "pages_crawled": len(pages),
        "pages_failed": len(failures),
        "sources": [page["source"] for page in pages],
    }


@app.post("/chatbot/{public_token}/knowledge/text")
def train_from_text(public_token: str, request: KnowledgeTextRequest):
    chatbot = get_chatbot_by_token(public_token)
    if not chatbot:
        raise HTTPException(status_code=404, detail="Chatbot not found")
    if len(request.text.strip()) < 20:
        raise HTTPException(status_code=422, detail="Knowledge text is too short")
    source = save_knowledge_source(
        chatbot["chatbot_id"],
        "text",
        request.title,
        request.text.strip(),
    )
    return {"trained": True, "source": source}


@app.get("/chatbot/{public_token}/knowledge")
def chatbot_knowledge(public_token: str):
    chatbot = get_chatbot_by_token(public_token)
    if not chatbot:
        raise HTTPException(status_code=404, detail="Chatbot not found")
    return {"chatbot": chatbot, "sources": get_knowledge_sources(chatbot["chatbot_id"])}


@app.get("/chatbot/{public_token}", response_class=HTMLResponse)
def hosted_chatbot(public_token: str):
    if not get_chatbot_by_token(public_token):
        raise HTTPException(status_code=404, detail="Chatbot not found")
    return chatbot_embed_markup(public_token)


@app.post("/unified-chat")
def unified_chat(request: UnifiedChatRequest):
    message = request.message.strip()
    surface = request.surface.lower().strip()

    if not message:
        raise HTTPException(status_code=422, detail="Message is required")
    if surface not in {"support", "website", "internal"}:
        raise HTTPException(status_code=422, detail="Unsupported surface")

    if surface == "internal":
        if not request.employee_id:
            raise HTTPException(status_code=401, detail="Employee SSO session is required")
        log_message(message, "employee", surface, employee_id=request.employee_id)
        if not request.customer_id:
            return {
                "reply": "Employee session verified. Provide a customer ID to perform a customer action.",
                "employee_id": request.employee_id,
            }
        customer = db_find_customer(request.customer_id)
        if not customer:
            raise HTTPException(status_code=404, detail="Customer not found")
        return {
            "reply": f"Customer {customer['name']} is ready for internal assistance.",
            "employee_id": request.employee_id,
            "customer": customer,
            "history": get_customer_history(customer["customer_id"]),
        }

    training_text = chatbot_training_text(request.chatbot_token)

    if (
        not request.phone_number
        and not request.customer_type
        and (is_support_issue(message) or is_account_action(message))
    ):
        return {
            "reply": (
                "I can help with that. Before I continue, are you an existing "
                "NetVeda customer? Please reply **existing** or **new**."
            ),
            "needs_customer_type": True,
        }

    if (
        not request.phone_number
        and request.customer_type == "existing"
    ):
        return {
            "reply": "Thanks. Please share your registered phone number so I can find your account and continue with the issue.",
            "needs_phone": True,
        }

    if (
        not request.phone_number
        and request.customer_type == "new"
        and not request.name
    ):
        return {
            "reply": "No problem. I’ll create a new customer record for this issue. Please share your name first.",
            "needs_name": True,
        }

    if (
        not request.phone_number
        and request.customer_type == "new"
        and request.name
    ):
        return {
            "reply": "Thanks. Please share a phone number so I can register the issue under your new customer account.",
            "needs_phone": True,
        }

    company_name = "NetVeda"
    if request.chatbot_token:
        bot_info = get_chatbot_by_token(request.chatbot_token)
        if bot_info:
            company_name = bot_info["name"]

    if (
        surface == "website"
        and not request.phone_number
        and is_plan_request(message)
        and not is_account_action(message)
    ):
        return {
            "reply": recommend_plan(message, company_name) + (
                "\n\nReply **199**, **299**, or **399** to compare a plan, "
                "or tell me if you want to buy, upgrade, or recharge."
            ),
            "guest_mode": True,
            "plan_options": [f"{company_name} 199", f"{company_name} 299", f"{company_name} 399"],
        }


    if not request.phone_number and not wants_ticket_or_contact(message):
        return {
            "reply": (
                generate_ai_response(
                    message,
                    request.conversation,
                    training_text,
                    use_global_knowledge=not bool(request.chatbot_token),
                )
                if client and not is_greeting(message)
                else conversational_reply(message)
            ),
            "guest_mode": True,
            "needs_phone": False,
            "offer_ticket": is_support_issue(message),
        }

    if surface == "website" and not request.phone_number and is_guest_eligible(message):
        return {"reply": recommend_plan(message, company_name), "guest_mode": True}

    if not request.phone_number:
        return {
            "reply": "Please share your registered phone number so I can connect this request to the correct account.",
            "needs_phone": True,
        }

    normalized_phone = normalize_phone(request.phone_number)
    if len(normalized_phone) < 10:
        raise HTTPException(status_code=422, detail="A valid phone number is required")

    customer = db_find_customer_by_phone(normalized_phone)
    created = False
    if not customer:
        if not request.name:
            return {
                "reply": "I could not find an existing account. Please provide your name to create one.",
                "needs_name": True,
            }
        customer = db_create_customer(
            name=request.name,
            phone=normalized_phone,
            email=request.email,
            product=request.product,
            source_surface=surface,
        )
        created = True

    ticket = db_create_ticket(
        customer_id=customer["customer_id"],
        query_text=message,
        category=classify_ticket(message),
        source_surface=surface,
    )
    save_customer_ticket_record(customer, ticket)
    reply_prefix = "Welcome to NetVeda" if created else f"Welcome back, {customer['name']}"
    return {
        "reply": f"{reply_prefix}. Your request has been registered. Ticket ID: {ticket['ticket_id']}",
        "customer": customer,
        "ticket": ticket,
        "created_customer": created,
    }

# =========================================================
# CHAT API
# =========================================================
@app.post("/chat")
async def chat(request: ChatRequest):
    message = request.message.strip()

    if not message:
        return {"reply": "Please enter a message."}

    # -----------------------------------------------------
    # ACTIVE SESSION
    # -----------------------------------------------------
    if customer_session.get("mode"):
        mode = customer_session.get("mode")

        if mode == "new_customer":
            return {"reply": handle_new_customer(message)}

        if mode == "existing_verification":
            return {"reply": handle_existing_verification(message)}

    # -----------------------------------------------------
    # EXISTING CUSTOMER REQUEST FIRST
    # -----------------------------------------------------
    if is_existing_customer_request(message):
        return {"reply": start_existing_customer_flow(message)}

    # -----------------------------------------------------
    # SUPPORT REQUEST
    # -----------------------------------------------------
    if is_support_issue(message):
        return {
            "reply": start_new_customer_flow(
                message,
                action="New Connection",
            )
        }

    # -----------------------------------------------------
    # ORDER REQUEST
    # -----------------------------------------------------
    if is_order_request(message):
        text = message.lower()
        action = "New Subscription"

        if "upgrade" in text:
            action = "Upgrade"
        elif "downgrade" in text:
            action = "Downgrade"
        elif "recharge" in text:
            action = "Recharge"
        elif "new connection" in text:
            action = "New Subscription"

        # If user clearly says they are already a customer, verify first.
        if "existing" in text or "already a customer" in text:
            return {"reply": start_existing_customer_flow(message)}

        return {
            "reply": start_new_customer_flow(message, action=action)
        }

    # -----------------------------------------------------
    # PLAN REQUEST
    # -----------------------------------------------------
    if is_plan_request(message):
        return {
            "reply": (
                recommend_plan(message)
                + "\n\n"
                "If you want to purchase or subscribe to a plan, "
                "tell me which plan you want."
            )
        }

    # -----------------------------------------------------
    # UNIVERSAL AI
    # -----------------------------------------------------
    return {"reply": generate_ai_response(message)}

# =========================================================
# HEALTH
# =========================================================
@app.get("/health")
def health():
    return {
        "status": "ok",
        "service": "NetVeda AI",
        "version": "10.0",
        "files": {
            "leads": str(LEADS_FILE),
            "support_tickets": str(SUPPORT_FILE),
            "orders_subscriptions": str(ORDERS_FILE),
        },
    }


@app.get("/employee/session")
def employee_session(employee_id: str | None = None):
    if not employee_id:
        raise HTTPException(status_code=401, detail="Employee SSO session is required")
    with get_connection() as connection:
        employee = connection.execute(
            "SELECT employee_id, name, role FROM employees WHERE employee_id = ?",
            (employee_id,),
        ).fetchone()
    if not employee:
        raise HTTPException(status_code=401, detail="Invalid employee session")
    return {"authenticated": True, "employee": dict(employee)}


@app.get("/analytics/summary")
def analytics_summary():
    with get_connection() as connection:
        customers = connection.execute("SELECT COUNT(*) FROM customers").fetchone()[0]
        tickets = connection.execute("SELECT COUNT(*) FROM tickets").fetchone()[0]
        open_tickets = connection.execute(
            "SELECT COUNT(*) FROM tickets WHERE status IN ('open', 'in-progress')"
        ).fetchone()[0]
        resolved_tickets = connection.execute(
            "SELECT COUNT(*) FROM tickets WHERE status = 'resolved'"
        ).fetchone()[0]
    resolution_rate = round((resolved_tickets / tickets) * 100, 2) if tickets else 0
    return {
        "customers": customers,
        "tickets": tickets,
        "open_tickets": open_tickets,
        "resolved_tickets": resolved_tickets,
        "resolution_rate": resolution_rate,
    }

# =========================================================
# TEST TICKET
# =========================================================
@app.post("/test-ticket")
def test_ticket():
    customer_id = generate_customer_id()
    ticket_id = create_support_ticket({
        "customer_id": customer_id,
        "issue": "Test support issue",
        "area": "Test Area",
        "priority": "Medium",
        "location": "Test Location",
    })

    return {
        "success": True,
        "customer_id": customer_id,
        "ticket_id": ticket_id,
        "status": "Open",
    }

# =========================================================
# MANUAL CREATE TICKET API
# =========================================================
@app.post("/create-ticket")
def create_ticket_api(request: TicketRequest):
    ticket_id = create_support_ticket({
        "customer_id": request.customer_id,
        "issue": request.issue,
        "area": request.area,
        "priority": request.priority,
        "location": request.location,
    })

    return {
        "success": True,
        "ticket_id": ticket_id,
        "customer_id": request.customer_id,
        "status": "Open",
        "priority": request.priority,
    }

# =========================================================
# ALL SUPPORT TICKETS
# =========================================================
@app.get("/support-tickets")
def support_tickets():
    create_support_file()
    wb = load_workbook(SUPPORT_FILE, data_only=True)
    ws = wb["Support Tickets"]
    tickets = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 7:
            continue
        tickets.append({
            "ticket_id": row[0],
            "customer_id": row[1],
            "issue": row[2],
            "area": row[3],
            "priority": row[4],
            "status": row[5],
            "location": row[6],
        })

    return {"tickets": tickets}

# =========================================================
# CUSTOMER TICKETS
# =========================================================
@app.get("/support-tickets/{customer_id}")
def customer_tickets(customer_id: str):
    return {
        "customer_id": customer_id,
        "tickets": get_customer_tickets(customer_id),
    }

# =========================================================
# ALL ORDERS & SUBSCRIPTIONS
# =========================================================
@app.get("/orders-subscriptions")
def orders_subscriptions():
    create_orders_file()
    wb = load_workbook(ORDERS_FILE, data_only=True)
    ws = wb["Orders & Subscriptions"]
    orders = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 9:
            continue
        orders.append({
            "order_id": row[0],
            "customer_id": row[1],
            "name": row[2],
            "phone": row[3],
            "plan": row[4],
            "amount": row[5],
            "action": row[6],
            "status": row[7],
            "date_time": row[8],
        })

    return {"orders": orders}

# =========================================================
# FRONTEND
# =========================================================
PUBLIC_PAGES = {
    "/about": "about.html",
    "/pricing": "pricing.html",
    "/blog": "blog.html",
    "/login": "login.html",
    "/app": "app.html",
}


for page_path, page_file in PUBLIC_PAGES.items():
    def serve_page(page_file=page_file):
        return FileResponse(FRONTEND_DIR / page_file)

    app.add_api_route(page_path, serve_page, methods=["GET"], include_in_schema=False)

if FRONTEND_DIR.exists():
    app.mount(
        "/",
        StaticFiles(directory=str(FRONTEND_DIR), html=True),
        name="frontend",
    )
    print("Frontend mounted successfully:", FRONTEND_DIR)
