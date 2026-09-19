from openpyxl import Workbook, load_workbook
from pathlib import Path
from datetime import datetime
import uuid


# ============================================================
# NETVEDA SUPPORT TICKET SYSTEM
# ============================================================

# Project paths
BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "Data"

# Excel file location
SUPPORT_TICKETS_FILE = DATA_DIR / "support_tickets.xlsx"


# ============================================================
# 1. CREATE EXCEL FILE
# ============================================================

def create_support_ticket_file():

    # Create Data folder if it does not exist
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    # If Excel file already exists, don't create it again
    if SUPPORT_TICKETS_FILE.exists():
        print("Support ticket file already exists.")
        return

    # Create workbook
    workbook = Workbook()

    # Get active sheet
    sheet = workbook.active

    # Rename sheet
    sheet.title = "Support Tickets"

    # Excel columns
    headers = [
        "Ticket ID",
        "Customer ID",
        "Issue",
        "Area",
        "Priority",
        "Status",
        "Location"
    ]

    # Add headers
    sheet.append(headers)

    # Save Excel file
    workbook.save(SUPPORT_TICKETS_FILE)

    print("support_tickets.xlsx created successfully.")


# ============================================================
# 2. GENERATE TICKET ID
# ============================================================

def generate_ticket_id():

    # Create random unique ID
    unique_id = uuid.uuid4().hex[:8].upper()

    return "TKT-" + unique_id


# ============================================================
# 3. CREATE NEW SUPPORT TICKET
# ============================================================

def create_support_ticket(
    customer_id,
    issue,
    area,
    priority="Medium",
    status="Open",
    location=""
):

    # Make sure Excel file exists
    create_support_ticket_file()

    try:

        # Open Excel file
        workbook = load_workbook(SUPPORT_TICKETS_FILE)

        # Select Support Tickets sheet
        sheet = workbook["Support Tickets"]

        # Generate Ticket ID
        ticket_id = generate_ticket_id()

        # Add ticket information
        sheet.append([
            ticket_id,
            customer_id,
            issue,
            area,
            priority,
            status,
            location
        ])

        # Save Excel file
        workbook.save(SUPPORT_TICKETS_FILE)

        print("----------------------------------------")
        print("SUPPORT TICKET CREATED")
        print("----------------------------------------")
        print("Ticket ID :", ticket_id)
        print("Customer ID:", customer_id)
        print("Issue      :", issue)
        print("Area       :", area)
        print("Priority   :", priority)
        print("Status     :", status)
        print("Location   :", location)
        print("----------------------------------------")

        return ticket_id

    except PermissionError:

        print("\nERROR:")
        print("Please close support_tickets.xlsx in Excel")
        print("and try again.")

        return None

    except Exception as error:

        print("\nERROR:", error)

        return None


# ============================================================
# 4. VIEW ALL SUPPORT TICKETS
# ============================================================

def view_all_tickets():

    create_support_ticket_file()

    try:

        workbook = load_workbook(SUPPORT_TICKETS_FILE)

        sheet = workbook["Support Tickets"]

        print("\n========================================")
        print("        NETVEDA SUPPORT TICKETS")
        print("========================================")

        # Read every row
        for row in sheet.iter_rows(values_only=True):

            print(row)

        print("========================================")

    except Exception as error:

        print("Error reading tickets:", error)


# ============================================================
# 5. FIND TICKET BY CUSTOMER ID
# ============================================================

def find_customer_tickets(customer_id):

    create_support_ticket_file()

    try:

        workbook = load_workbook(SUPPORT_TICKETS_FILE)

        sheet = workbook["Support Tickets"]

        print("\n========================================")
        print("CUSTOMER TICKETS")
        print("========================================")

        found = False

        # Skip header row
        for row in sheet.iter_rows(min_row=2, values_only=True):

            ticket_id = row[0]
            row_customer_id = row[1]
            issue = row[2]
            area = row[3]
            priority = row[4]
            status = row[5]
            location = row[6]

            if row_customer_id == customer_id:

                found = True

                print("\nTicket ID :", ticket_id)
                print("Issue     :", issue)
                print("Area      :", area)
                print("Priority  :", priority)
                print("Status    :", status)
                print("Location  :", location)

        if not found:

            print("No tickets found for", customer_id)

        print("========================================")

    except Exception as error:

        print("Error:", error)


# ============================================================
# 6. UPDATE TICKET STATUS
# ============================================================

def update_ticket_status(ticket_id, new_status):

    create_support_ticket_file()

    try:

        workbook = load_workbook(SUPPORT_TICKETS_FILE)

        sheet = workbook["Support Tickets"]

        found = False

        # Search tickets
        for row in sheet.iter_rows(min_row=2):

            current_ticket_id = row[0].value

            if current_ticket_id == ticket_id:

                # Status is column 6
                row[5].value = new_status

                found = True

                break

        if found:

            workbook.save(SUPPORT_TICKETS_FILE)

            print("Ticket status updated successfully.")

        else:

            print("Ticket not found.")

    except PermissionError:

        print("Please close support_tickets.xlsx first.")

    except Exception as error:

        print("Error:", error)


# ============================================================
# 7. MAIN TEST PROGRAM
# ============================================================

if __name__ == "__main__":

    print("\n========================================")
    print("       NETVEDA SUPPORT SYSTEM")
    print("========================================")

    # Create Excel file
    create_support_ticket_file()

    # --------------------------------------------------------
    # TEST 1: Create a support ticket
    # --------------------------------------------------------

    ticket = create_support_ticket(
        customer_id="NV1001",
        issue="Slow Internet",
        area="Ludhiana",
        priority="Medium",
        status="Open",
        location="Ludhiana, Punjab"
    )

    # --------------------------------------------------------
    # TEST 2: View all tickets
    # --------------------------------------------------------

    view_all_tickets()

    # --------------------------------------------------------
    # TEST 3: Find tickets of customer
    # --------------------------------------------------------

    find_customer_tickets("NV1001")

    # --------------------------------------------------------
    # TEST 4: Update ticket status
    # --------------------------------------------------------

    if ticket:

        update_ticket_status(
            ticket,
            "In Progress"
        )

    print("\nSupport ticket system test completed.")